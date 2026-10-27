import os
import tempfile
import unittest
import zipfile

from tests.helpers import CASE1, PHOTO, load_case

try:
    import docx  # noqa: F401
    from build_docx import TEMPLATES, build_cv, build_letter
    HAVE_DOCX = True
except ImportError:
    HAVE_DOCX = False


@unittest.skipUnless(HAVE_DOCX, 'python-docx not installed')
class BuildDocxTest(unittest.TestCase):
    def setUp(self):
        _, self.output, _ = load_case(CASE1)
        self.tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self.tmp.cleanup()

    def path(self, name):
        return os.path.join(self.tmp.name, name)

    def text(self, path):
        import docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for cell in t._cells:
                parts += [p.text for p in cell.paragraphs]
        return '\n'.join(parts)

    def test_every_template_builds_with_its_font(self):
        for name, spec in TEMPLATES.items():
            p = self.path(f'{name}.docx')
            build_cv(self.output['cv'], p, template=name)
            styles = zipfile.ZipFile(p).read('word/styles.xml').decode()
            self.assertIn(spec['font'], styles, name)
            self.assertIn('Harbourline Group', self.text(p), name)

    def test_page_sizes(self):
        import docx
        build_cv(self.output['cv'], self.path('us.docx'), page='letter')
        build_cv(self.output['cv'], self.path('eu.docx'), page='a4')
        self.assertEqual(round(docx.Document(self.path('us.docx')).sections[0].page_width.mm), 216)
        self.assertEqual(round(docx.Document(self.path('eu.docx')).sections[0].page_width.mm), 210)

    def test_regional_extras(self):
        p = self.path('uae.docx')
        build_cv(self.output['cv'], p, template='modern', photo=PHOTO,
                 header_lines=['Visa status: Employment visa (transferable)'],
                 details=[('Nationality', 'Irish'), ('Notice period', '30 days')], languages='English (native), Arabic (A2)',
                 declaration=True, place='Dubai', date='27 October 2026')
        names = zipfile.ZipFile(p).namelist()
        self.assertTrue(any(n.startswith('word/media/') for n in names), names)
        text = self.text(p)
        for s in ('Visa status: Employment visa (transferable)', 'PERSONAL DETAILS', 'Nationality: Irish', 'LANGUAGES', 'DECLARATION', 'Dubai, 27 October 2026'):
            self.assertIn(s, text)

    def test_placeholders_highlighted(self):
        p = self.path('cv.docx')
        build_cv(self.output['cv'], p)
        xml = zipfile.ZipFile(p).read('word/document.xml').decode()
        self.assertIn('w:highlight w:val="yellow"', xml)

    def test_cover_letter(self):
        p = self.path('letter.docx')
        build_letter(self.output['cover_letter'], p, name='Jordan Byrne', date='27 October 2026')
        text = self.text(p)
        self.assertIn('Dear Hiring Manager', text)
        self.assertIn('Kind regards', text)

    def test_cli_builds_only_the_stage_that_ran(self):
        import json
        import subprocess
        import sys
        script = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'shared', 'scripts', 'build_docx.py')
        stages = {
            'cv-only': ({'analysis': self.output['analysis'], 'cv': self.output['cv']}, ['jordan-byrne-cv.docx']),
            'letter-only': ({'analysis': self.output['analysis'], 'cover_letter': self.output['cover_letter']}, ['jordan-byrne-cover-letter.docx']),
        }
        for label, (data, expected) in stages.items():
            outdir = self.path(label)
            src = self.path(f'{label}.json')
            json.dump(data, open(src, 'w'))
            subprocess.run([sys.executable, script, src, '--outdir', outdir, '--name', 'Jordan Byrne'], check=True, capture_output=True)
            self.assertEqual(sorted(os.listdir(outdir)), expected, label)


if __name__ == '__main__':
    unittest.main()
