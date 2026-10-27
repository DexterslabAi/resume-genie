#!/usr/bin/env python3
"""Build Word (.docx) files for a tailored CV and cover letter in Resume Genie's templates.

Templates: classic (centred, safest for ATS), modern (violet accent), executive (Cambria,
double rules), compact (tighter). Placeholders like [N] or [X%] are highlighted in yellow
so the candidate can't miss them.

Usage:
  python build_docx.py OUTPUT.json [--template classic] [--page a4|letter] [--outdir .]
      [--photo photo.jpg] [--header "Visa status: Employment visa (transferable)"]
      [--detail "Nationality=Irish"] [--languages "English (native), Arabic (A2)"]
      [--declaration] [--signature] [--place Dubai] [--date "27 October 2026"]

OUTPUT.json is a Resume Genie output (see references/output-schema.json).
Requires python-docx (pip install python-docx).

Author: Rahul Kashyap. MIT licence.
"""
import argparse
import datetime
import json
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

TEMPLATES = {  # font, heading colour, centred header, body size, name size
    'classic':   dict(font='Calibri', accent='1F2937', center=True,  body=10.5, name=20, rule='single'),
    'modern':    dict(font='Calibri', accent='5B3DF5', center=False, body=10.5, name=22, rule='thick'),
    'executive': dict(font='Cambria', accent='1B1446', center=True,  body=10.5, name=22, rule='double'),
    'compact':   dict(font='Calibri', accent='374151', center=False, body=9.5,  name=17, rule='single'),
}

def _rule(p, kind, color):
    pPr = p._p.get_or_add_pPr(); b = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'double' if kind == 'double' else 'single'); bot.set(qn('w:sz'), '14' if kind == 'thick' else '6')
    bot.set(qn('w:space'), '2'); bot.set(qn('w:color'), color); b.append(bot); pPr.append(b)

def _runs(p, text, size, bold=False, italic=False, color=None):
    for part in re.split(r'(\[[^\]]+\])', str(text or '')):
        if not part: continue
        r = p.add_run(part); r.font.size = Pt(size); r.bold = bold; r.italic = italic
        if color: r.font.color.rgb = RGBColor.from_string(color)
        if re.fullmatch(r'\[[^\]]+\]', part): r.font.highlight_color = WD_COLOR_INDEX.YELLOW  # placeholder to fill in

def build_cv(cv, path, template='classic', page='a4', photo=None, header_lines=(), details=(), languages='', declaration=False, signature=False, place='', date=''):
    T = TEMPLATES.get(template, TEMPLATES['classic'])
    doc = Document()
    st = doc.styles['Normal']; st.font.name = T['font']; st.font.size = Pt(T['body'])
    st.element.rPr.rFonts.set(qn('w:eastAsia'), T['font'])
    sec = doc.sections[0]
    sec.page_width, sec.page_height = (Mm(215.9), Mm(279.4)) if page == 'letter' else (Mm(210), Mm(297))
    sec.left_margin = sec.right_margin = Cm(1.76); sec.top_margin = sec.bottom_margin = Cm(1.6)
    align = WD_ALIGN_PARAGRAPH.LEFT if (photo or not T['center']) else WD_ALIGN_PARAGRAPH.CENTER
    head = doc
    if photo:  # two-column header: text left, photo right
        tbl = doc.add_table(rows=1, cols=2); tbl.autofit = False
        tbl.columns[0].width = Cm(14); tbl.columns[1].width = Cm(3.4)
        head = tbl.cell(0, 0); pic = tbl.cell(0, 1).paragraphs[0]; pic.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        pic.add_run().add_picture(photo, width=Cm(3), height=Cm(3.75))
    def hp(text, size, bold=False, color=None):
        p = head.add_paragraph() if photo or head is not doc else doc.add_paragraph()
        p.alignment = align; p.paragraph_format.space_after = Pt(2); _runs(p, text, size, bold, color=color); return p
    if photo: head.paragraphs[0]._p.getparent().remove(head.paragraphs[0]._p)
    hp(cv.get('name', ''), T['name'], True, T['accent'] if template == 'modern' else None)
    if cv.get('headline'): hp(cv['headline'], T['body'] + 0.5, color='374151')
    if cv.get('contact'): hp('  |  '.join(cv['contact']), T['body'] - 1, color='4B5563')
    for line in header_lines: hp(line, T['body'] - 1, True, '4B5563')
    def heading(text):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10 if template != 'compact' else 7); p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True; _runs(p, text.upper(), T['body'], True, color=T['accent']); _rule(p, T['rule'], T['accent'])
    def para(text, bold_prefix=''):
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        if bold_prefix: _runs(p, bold_prefix, T['body'], True)
        _runs(p, text, T['body']); return p
    def bullet(text):
        p = doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after = Pt(1); _runs(p, text, T['body'])
    if cv.get('summary'): heading('Professional Summary'); para(cv['summary'])
    if cv.get('skills'):
        heading('Key Skills')
        for g in cv['skills']: para(', '.join(g.get('items', [])), f"{g['group']}: " if g.get('group') else '')
    if cv.get('experience'):
        heading('Professional Experience')
        width = sec.page_width - sec.left_margin - sec.right_margin
        for job in cv['experience']:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(0); p.paragraph_format.keep_with_next = True
            p.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
            _runs(p, job.get('title', ''), T['body'] + 0.5, True); _runs(p, '\t' + job.get('dates', ''), T['body'] - 0.5, color='4B5563')
            q = doc.add_paragraph(); q.paragraph_format.space_after = Pt(2)
            _runs(q, ', '.join(x for x in [job.get('company'), job.get('location')] if x), T['body'], italic=True, color=T['accent'] if template == 'modern' else '374151')
            for b in job.get('bullets', []): bullet(b['text'] if isinstance(b, dict) else b)
    if cv.get('projects'):
        heading('Projects')
        for pr in cv['projects']:
            para(' | '.join(x for x in [pr.get('org'), pr.get('dates')] if x), f"{pr.get('title', '')}  ")
            if pr.get('details'): bullet(pr['details'])
    if cv.get('education'):
        heading('Education')
        for e in cv['education']:
            para(', '.join(x for x in [e.get('institution'), e.get('dates')] if x), f"{e.get('degree', '')}, " if e.get('institution') or e.get('dates') else e.get('degree', ''))
            if e.get('details'): bullet(e['details'])
    if cv.get('certifications'):
        heading('Certifications')
        for c in cv['certifications']: bullet(c)
    if languages: heading('Languages'); para(languages)
    if details:
        heading('Personal Details')
        for label, value in details: para(value, f'{label}: ')
    if declaration:
        heading('Declaration'); para('I hereby declare that the information given above is true to the best of my knowledge and belief.')
    if declaration or signature:
        p = para(f"{place or 'Place'}, {date or 'Date'}"); p.paragraph_format.space_before = Pt(18)
        para('_______________________   ' + cv.get('name', ''))
    doc.save(path)

def build_letter(letter, path, name='', template='classic', page='a4', date=''):
    T = TEMPLATES.get(template, TEMPLATES['classic'])
    doc = Document()
    st = doc.styles['Normal']; st.font.name = T['font']; st.font.size = Pt(11)
    sec = doc.sections[0]; sec.page_width, sec.page_height = (Mm(215.9), Mm(279.4)) if page == 'letter' else (Mm(210), Mm(297))
    for text, size, bold in [(name, 14, True), (date, 11, False), (letter.get('greeting', ''), 11, False)] + [(p, 11, False) for p in letter.get('paragraphs', [])] + [(letter.get('sign_off', ''), 11, False)]:
        if text: p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(10); _runs(p, text, size, bold)
    doc.save(path)

def slug(name):
    return re.sub(r'[^a-z0-9]+', '-', str(name or 'cv').lower()).strip('-') or 'cv'


def main():
    ap = argparse.ArgumentParser(description='Build CV and cover letter .docx files from a Resume Genie output.')
    ap.add_argument('output_json')
    ap.add_argument('--template', default='classic', choices=sorted(TEMPLATES))
    ap.add_argument('--page', default='a4', choices=['a4', 'letter'])
    ap.add_argument('--outdir', default='.')
    ap.add_argument('--photo')
    ap.add_argument('--header', action='append', default=[], help='extra header line, e.g. visa status')
    ap.add_argument('--detail', action='append', default=[], help='personal detail as Label=Value')
    ap.add_argument('--languages', default='')
    ap.add_argument('--declaration', action='store_true')
    ap.add_argument('--signature', action='store_true')
    ap.add_argument('--place', default='')
    today = datetime.date.today()
    ap.add_argument('--date', default=f'{today.day} {today:%B %Y}')
    a = ap.parse_args()

    data = json.load(open(a.output_json, encoding='utf-8'))
    cv = data['cv']
    details = [tuple(d.split('=', 1)) for d in a.detail if '=' in d]
    os.makedirs(a.outdir, exist_ok=True)
    cv_path = os.path.join(a.outdir, f"{slug(cv.get('name'))}-cv.docx")
    build_cv(cv, cv_path, template=a.template, page=a.page, photo=a.photo, header_lines=a.header, details=details,
             languages=a.languages, declaration=a.declaration, signature=a.signature, place=a.place, date=a.date)
    print('Wrote', cv_path)
    if data.get('cover_letter'):
        letter_path = os.path.join(a.outdir, f"{slug(cv.get('name'))}-cover-letter.docx")
        build_letter(data['cover_letter'], letter_path, name=cv.get('name', ''), template=a.template, page=a.page, date=a.date)
        print('Wrote', letter_path)


if __name__ == '__main__':
    main()
